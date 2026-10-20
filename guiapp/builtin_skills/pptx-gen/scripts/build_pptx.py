"""Build a designed .pptx deck from a JSON outline using python-pptx.

The deck uses a real visual system (cover, accent rail, cards, page footer),
matching the native office write_pptx renderer. Input JSON:

    {"title": "...", "subtitle": "...", "theme": "clean|magazine|eink|dashboard|retro|sketch|whiteboard|handmade|research|mckinsey|official|course",
     "slides": [{"title": "...", "kicker": "...",
                 "layout": "auto|bullets|cards|timeline|comparison|section|agenda|kpi|quote|closing",
                 "bullets": ["..."], "notes": "...",
                 "images": [{"path": "..."}],
                 "charts": [{"chart_type": "line", "title": "...",
                             "categories": ["..."],
                             "series": [{"name": "...", "values": [1]}]}]}]}

Usage:
    python build_pptx.py --input outline.json --output deck.pptx
    python build_pptx.py --input outline.json --output ./out_dir
    echo '{...}' | python build_pptx.py --output deck.pptx
"""
import argparse
import json
import os
import re
import sys

_ILLEGAL_XML_CHARS_RE = re.compile(
    r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f'
    r'\ud800-\udfff\ufdd0-\ufdef\ufffe\uffff]'
)

_DEFAULT_FONT = "Microsoft YaHei"

# Keep ids, aliases, keywords, colors, and motif aligned with corelib/pptx/write_style.go.
_STYLES = [
    {"id": "clean", "motif": "report", "aliases": ["clean", "business", "商务", "商务汇报", "汇报", "清爽专业", "clean professional"],
     "keywords": ["经营汇报", "工作总结", "项目复盘", "晋升述职", "季度汇报", "客户提案", "工作复盘", "清爽专业", "商务汇报", "quarterly review", "business review", "client proposal"],
     "colors": {"paper": "FFFFFF", "ink": "0F172A", "navy": "1E3A8A", "navy2": "0F766E", "accent": "2563EB", "gold": "F59E0B", "white": "FFFFFF", "mute": "64748B", "slate": "334155", "card": "F8FAFC", "on_dark": "F8FAFC", "on_dark_mute": "CBD5E1"}},
    {"id": "magazine", "motif": "split", "aliases": ["magazine", "launch", "发布", "发布路演", "路演", "modern", "现代", "现代品牌", "品牌", "创意杂志", "editorial magazine"],
     "keywords": ["品牌发布", "创意提案", "设计作品集", "时尚发布", "产品发布", "产品介绍", "创意杂志", "发布路演", "product launch", "pitch deck", "fundraising", "brand deck"],
     "colors": {"paper": "FFFFFF", "ink": "111111", "navy": "111111", "navy2": "2A2A2A", "accent": "FF006E", "gold": "FFED00", "white": "FFFFFF", "mute": "6B7280", "slate": "374151", "card": "F4F4F5", "on_dark": "FFFFFF", "on_dark_mute": "E5E7EB"}},
    {"id": "eink", "motif": "ink", "aliases": ["eink", "minimal", "极简", "ceremony", "年会", "典礼", "e-ink", "电子墨水", "e-ink magazine"],
     "keywords": ["主题演讲", "观点分享", "个人介绍", "年会致辞", "颁奖典礼", "电子墨水", "极简风格", "demo day", "personal introduction"],
     "colors": {"paper": "F7F4EA", "ink": "111111", "navy": "1E2A78", "navy2": "163B2F", "accent": "C2410C", "gold": "D9C29E", "white": "FFFCF5", "mute": "6B6560", "slate": "3F3A36", "card": "FBF8F1", "on_dark": "F7F4EA", "on_dark_mute": "D6D0C4"}},
    {"id": "dashboard", "motif": "dashboard", "aliases": ["dashboard", "数据仪表盘", "看板", "data dashboard"],
     "keywords": ["数据看板", "运营分析", "业绩展示", "商业智能", "KPI汇报", "数据仪表盘", "看板", "analytics dashboard"],
     "colors": {"paper": "F5F8FC", "ink": "0F172A", "navy": "1D4ED8", "navy2": "0891B2", "accent": "2563EB", "gold": "10B981", "white": "FFFFFF", "mute": "64748B", "slate": "334155", "card": "FFFFFF", "on_dark": "F8FAFC", "on_dark_mute": "C7D2FE"}},
    {"id": "retro", "motif": "poster", "aliases": ["retro", "复古", "扁平插画", "复古扁平", "retro flat"],
     "keywords": ["品牌故事", "生活方式", "旅游介绍", "文化创意", "复古插画", "复古扁平", "复古风格", "brand story"],
     "colors": {"paper": "F5F3E8", "ink": "34495E", "navy": "2C3E50", "navy2": "6C7A89", "accent": "FF6B6B", "gold": "F9CA24", "white": "FFFFFF", "mute": "6C7A89", "slate": "34495E", "card": "FFFDF6", "on_dark": "2C3E50", "on_dark_mute": "95A5A6"}},
    {"id": "sketch", "motif": "sketch", "aliases": ["sketch", "tech", "技术", "技术分享", "手绘技术", "sketch explainer"],
     "keywords": ["技术分享", "系统架构", "技术架构", "概念解释", "开源项目", "技术方案", "手绘技术", "engineering talk", "system architecture"],
     "colors": {"paper": "FCFBF7", "ink": "2F3437", "navy": "2F3437", "navy2": "7A8490", "accent": "BFD7F1", "gold": "F4C7B8", "white": "FFFFFF", "mute": "8A9096", "slate": "4A5158", "card": "FFFFFF", "on_dark": "2F3437", "on_dark_mute": "8A9096"}},
    {"id": "whiteboard", "motif": "board", "aliases": ["whiteboard", "白板", "手绘白板"],
     "keywords": ["头脑风暴", "内部研讨", "白板讲解", "概念梳理", "手绘白板", "白板风格", "brainstorm"],
     "colors": {"paper": "FAFAF5", "ink": "1A1A1A", "navy": "1A1A1A", "navy2": "3498DB", "accent": "E74C3C", "gold": "F39C12", "white": "FFFFFF", "mute": "6B7280", "slate": "374151", "card": "FFF9C4", "on_dark": "FFFFFF", "on_dark_mute": "D1D5DB"}},
    {"id": "handmade", "motif": "craft", "aliases": ["handmade", "warm", "温馨", "温馨纪念", "温暖手工"],
     "keywords": ["生日纪念", "生日", "温馨纪念", "儿童教育", "亲子活动", "社区公益", "温暖手工", "birthday", "wedding", "family gathering"],
     "colors": {"paper": "F5F1E8", "ink": "5C4033", "navy": "5C4033", "navy2": "A67C52", "accent": "F5C4B8", "gold": "B8D4A8", "white": "FFFCF8", "mute": "8A7568", "slate": "6B5344", "card": "FFFCF7", "on_dark": "5C4033", "on_dark_mute": "A67C52"}},
    {"id": "research", "motif": "defense", "aliases": ["research", "academic", "学术", "学术答辩", "科研答辩", "research defense"],
     "keywords": ["论文答辩", "开题答辩", "毕业答辩", "课题申报", "基金申请", "文献综述", "科研答辩", "学术答辩", "答辩", "thesis defense", "dissertation"],
     "colors": {"paper": "FFFFFF", "ink": "111111", "navy": "003F8F", "navy2": "0B5CAD", "accent": "B5121B", "gold": "F3DFA2", "white": "FFFFFF", "mute": "4B5563", "slate": "333333", "card": "F3F6FA", "on_dark": "FFFFFF", "on_dark_mute": "D6E4F0"}},
    {"id": "mckinsey", "motif": "consulting", "aliases": ["mckinsey", "consulting", "咨询", "咨询顾问", "麦肯锡", "executive", "高管", "高管汇报"],
     "keywords": ["战略咨询", "商业分析", "行业分析", "尽职调查", "竞对分析", "增长战略", "战略汇报", "董事会汇报", "麦肯锡", "咨询风格", "due diligence", "strategy consulting"],
     "colors": {"paper": "F6F8FA", "ink": "243447", "navy": "243447", "navy2": "5B7C99", "accent": "5B7C99", "gold": "8AA0B5", "white": "FFFFFF", "mute": "6B7280", "slate": "475569", "card": "FFFFFF", "on_dark": "243447", "on_dark_mute": "6B7280"}},
    {"id": "official", "motif": "official", "aliases": ["official", "党政", "党政红", "official red"],
     "keywords": ["党建学习", "政策宣讲", "工作部署", "党政汇报", "国企汇报", "党政红", "党建汇报"],
     "colors": {"paper": "FFF9F2", "ink": "262626", "navy": "9E1530", "navy2": "C41E3A", "accent": "C41E3A", "gold": "D6A84B", "white": "FFFFFF", "mute": "6B6560", "slate": "4A4A4A", "card": "FFFCF8", "on_dark": "FFF9F2", "on_dark_mute": "F3E6C8"}},
    {"id": "course", "motif": "lesson", "aliases": ["course", "education", "教学", "课件", "教学课件", "lesson deck"],
     "keywords": ["教学课件", "培训课程", "培训课件", "课堂教案", "新员工培训", "课件", "教案", "lesson plan", "classroom"],
     "colors": {"paper": "FFFBF5", "ink": "1F2937", "navy": "1E3A5F", "navy2": "0F766E", "accent": "F59E0B", "gold": "FDE68A", "white": "FFFFFF", "mute": "6B7280", "slate": "374151", "card": "FFFCF7", "on_dark": "FFFBF5", "on_dark_mute": "E5E7EB"}},
]


def _clean_xml_text(text):
    if not text:
        return ""
    return _ILLEGAL_XML_CHARS_RE.sub("", str(text)).strip()


def _rgb(hex_code):
    from pptx.dml.color import RGBColor
    hex_code = hex_code.strip().lstrip("#")
    return RGBColor(int(hex_code[0:2], 16), int(hex_code[2:4], 16), int(hex_code[4:6], 16))


def _set_run_font(run, size=None, bold=None, color=None):
    from pptx.oxml.ns import qn
    run.font.name = _DEFAULT_FONT
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:latin"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = r_pr.makeelement(qn(tag), {})
            r_pr.append(node)
        node.set("typeface", _DEFAULT_FONT)


def _load_outline(args):
    if args.input:
        if not os.path.isfile(args.input):
            print(f"ERROR: Input JSON file not found: {args.input}", file=sys.stderr)
            sys.exit(1)
        with open(args.input, "r", encoding="utf-8") as handle:
            raw = handle.read()
    else:
        raw = sys.stdin.read()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        print(f"ERROR: Invalid outline JSON: {exc}", file=sys.stderr)
        sys.exit(1)
    if not isinstance(data, dict):
        print("ERROR: Outline JSON must be an object with 'title' and 'slides'.", file=sys.stderr)
        sys.exit(1)
    slides = data.get("slides")
    if not isinstance(slides, list) or not slides:
        print("ERROR: Outline JSON must contain a non-empty 'slides' array.", file=sys.stderr)
        sys.exit(1)
    return data


def _resolve_output_path(output, title):
    if os.path.isdir(output):
        safe = re.sub(r'[\\/:*?"<>|]+', "_", (title or "presentation").strip()) or "presentation"
        output = os.path.join(output, f"{safe}.pptx")
    elif not output.lower().endswith(".pptx"):
        output += ".pptx"
    parent = os.path.dirname(os.path.abspath(output))
    if parent:
        os.makedirs(parent, exist_ok=True)
    return output


def _lookup_style(name):
    key = str(name or "").strip().lower()
    if key.startswith("theme="):
        key = key[len("theme="):].strip()
    if key in ("", "auto", "自动", "default", "默认"):
        return None
    for style in _STYLES:
        names = [style["id"], *style["aliases"]]
        if key in [item.lower() for item in names]:
            return style
    return None


def _match_style(hint):
    hint = str(hint or "").lower()
    m = re.search(r"theme=([a-z0-9_-]+)", hint)
    if m:
        found = _lookup_style(m.group(1))
        if found:
            return found
    best = None
    best_len = 0
    best_score = 0
    for style in _STYLES:
        score = 0
        longest = 0
        for keyword in style["keywords"]:
            if keyword.lower() in hint:
                n = len(keyword)
                score += n
                longest = max(longest, n)
        if score and (longest > best_len or (longest == best_len and score > best_score)):
            best = style
            best_len = longest
            best_score = score
    return best


def _style_hint(outline):
    parts = [outline.get("purpose"), outline.get("title"), outline.get("subtitle")]
    for item in outline.get("slides") or []:
        if isinstance(item, dict):
            parts.extend([item.get("title"), item.get("kicker")])
    return " ".join(_clean_xml_text(part) for part in parts if part)


def _resolve_theme(name, hint):
    found = _lookup_style(name)
    if found:
        return found
    text = hint
    if str(name or "").strip().lower() not in ("", "auto", "自动", "default", "默认"):
        text = f"{name} {hint}"
    return _match_style(text) or _STYLES[0]


def _texts(bullets):
    if isinstance(bullets, str):
        bullets = [bullets]
    out = []
    for item in bullets or []:
        text = _clean_xml_text(item)
        if text:
            out.append(text)
    return out


def _layout(item):
    raw = str(item.get("layout") or "auto").strip().lower()
    aliases = {
        "": "auto", "自动": "auto",
        "bullet": "bullets", "list": "bullets", "要点": "bullets", "列表": "bullets",
        "card": "cards", "卡片": "cards", "grid": "cards", "网格": "cards",
        "process": "timeline", "steps": "timeline", "flow": "timeline", "roadmap": "timeline", "时间线": "timeline", "流程": "timeline", "步骤": "timeline", "路线图": "timeline", "阶段": "timeline",
        "compare": "comparison", "versus": "comparison", "vs": "comparison", "split": "comparison", "two_column": "comparison", "two-column": "comparison", "columns": "comparison", "对比": "comparison", "对照": "comparison", "双栏": "comparison", "比较": "comparison",
        "divider": "section", "分节": "section", "章节": "section",
        "toc": "agenda", "目录": "agenda", "议程": "agenda",
        "metrics": "kpi", "指标": "kpi",
        "引述": "quote", "金句": "quote",
        "close": "closing", "end": "closing", "结尾": "closing", "结束": "closing",
    }
    kind = aliases.get(raw, raw if raw in ("auto", "bullets", "cards", "timeline", "comparison", "section", "agenda", "kpi", "quote", "closing") else "auto")
    texts = _texts(item.get("bullets"))
    has_media = bool(item.get("images") or item.get("charts"))
    if has_media:
        return "bullets"
    if kind == "cards" and (not texts or len(texts) > 6):
        return "bullets"
    if kind in ("timeline", "comparison") and (len(texts) < 2 or len(texts) > 6):
        return "bullets"
    if kind == "agenda" and (not texts or len(texts) > 8):
        return "bullets"
    if kind == "kpi" and (not texts or len(texts) > 6):
        return "bullets"
    if kind != "auto":
        return kind
    if not texts:
        return "section"
    header = f"{item.get('title') or ''} {item.get('kicker') or ''}".lower()
    if 3 <= len(texts) <= 6 and all(len(t) <= 64 for t in texts):
        if any(kw in header for kw in ("流程", "时间线", "路线图", "步骤", "演进", "阶段", "路径", "roadmap", "timeline", "workflow")):
            return "timeline"
    if len(texts) in (2, 4, 6) and all(len(t) <= 64 for t in texts):
        if any(kw in header for kw in ("对比", "对照", "比较", "优劣", "差异", " vs ", "versus")):
            return "comparison"
    if 2 <= len(texts) <= 4 and all(len(text) <= 52 for text in texts):
        return "cards"
    return "bullets"


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if "blank" in (layout.name or "").lower():
            return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


def _no_line(shape):
    shape.line.fill.background()


def _rect(slide, inches, x, y, w, h, fill, rounded=False):
    from pptx.enum.shapes import MSO_SHAPE
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, inches(x), inches(y), inches(w), inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    _no_line(shape)
    if rounded:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def _tb(slide, inches, pt, x, y, w, h, text, size, color, bold=False, align="left"):
    from pptx.enum.text import PP_ALIGN
    text = _clean_xml_text(text)
    if not text:
        return None
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
    run = paragraph.add_run()
    run.text = text
    _set_run_font(run, size=pt(size), bold=bold, color=color)
    return box


def _lines(slide, inches, pt, x, y, w, h, lines, size, color, space=8):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    lines = _texts(lines)
    if not lines:
        return
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(space)
        run = paragraph.add_run()
        run.text = line
        _set_run_font(run, size=pt(size), color=color)


def _page_mark(slide, inches, pt, page, total, color):
    _tb(slide, inches, pt, 11.05, 7.08, 1.8, 0.28, f"{page:02d}  /  {total:02d}", 12, color, align="right")


def _light_chrome(slide, inches, pt, theme, kicker, title, footer, page, total, motif="report"):
    _rect(slide, inches, 0, 0, 13.333, 7.5, theme["paper"])
    if motif == "official":
        _rect(slide, inches, 0, 0, 13.333, 0.12, theme["accent"])
        _rect(slide, inches, 0, 0.12, 13.333, 0.03, theme["gold"])
    elif motif == "lesson":
        _rect(slide, inches, 0, 0, 13.333, 0.16, theme["navy"])
        _rect(slide, inches, 0, 0.16, 13.333, 0.04, theme["accent"])
    elif motif == "split":
        _rect(slide, inches, 0, 0, 0.12, 7.5, theme["accent"])
    elif motif == "consulting":
        _rect(slide, inches, 0.55, 1.02, 12.3, 0.012, theme["accent"])
    elif motif == "defense":
        _rect(slide, inches, 0.55, 1.02, 1.7, 0.04, theme["navy"])
        _rect(slide, inches, 2.35, 1.02, 0.4, 0.04, theme["accent"])
    elif motif not in ("ink", "sketch", "dashboard", "board"):
        _rect(slide, inches, 0, 0, 0.14, 7.5, theme["accent"])
    title_y = 0.32
    title_size = 28
    if kicker:
        _tb(slide, inches, pt, 0.55, 0.18, 12.0, 0.26, kicker, 12, theme["accent"], bold=True)
        title_y = 0.46
        title_size = 26
    if title:
        _tb(slide, inches, pt, 0.55, title_y, 12.0, 0.52, title, title_size, theme["ink"], bold=True)
    if footer:
        _tb(slide, inches, pt, 0.55, 7.08, 9.3, 0.28, footer, 11, theme["mute"])
    _page_mark(slide, inches, pt, page, total, theme["mute"])
    return 1.42 if kicker else 1.2


def _paint_cover(slide, inches, pt, motif, theme, title, subtitle, total):
    if motif == "split":
        _rect(slide, inches, 0, 0, 5.55, 7.5, theme["navy"])
        _rect(slide, inches, 5.55, 0, 7.783, 7.5, theme["paper"])
        _rect(slide, inches, 5.55, 0, 0.12, 7.5, theme["accent"])
        _tb(slide, inches, pt, 0.45, 2.05, 4.7, 2.2, title, 36, theme["white"], bold=True)
        if subtitle:
            _tb(slide, inches, pt, 6.15, 2.3, 6.5, 1.6, subtitle, 18, theme["ink"])
        return
    if motif == "official":
        _rect(slide, inches, 0, 0, 13.333, 7.5, theme["paper"])
        _rect(slide, inches, 0, 0, 5.15, 7.5, theme["accent"])
        _rect(slide, inches, 5.15, 0, 0.045, 7.5, theme["gold"])
        _tb(slide, inches, pt, 0.45, 2.15, 4.3, 2.2, title, 32, theme["white"], bold=True)
        if subtitle:
            _tb(slide, inches, pt, 5.7, 2.4, 7.0, 1.6, subtitle, 18, theme["ink"])
        return
    if motif == "lesson":
        _rect(slide, inches, 0, 0, 13.333, 7.5, theme["paper"])
        _rect(slide, inches, 0, 0, 13.333, 2.45, theme["navy"])
        _rect(slide, inches, 0, 2.45, 13.333, 0.08, theme["accent"])
        _tb(slide, inches, pt, 0.7, 0.85, 11.6, 1.15, title, 32, theme["white"], bold=True)
        if subtitle:
            _tb(slide, inches, pt, 0.7, 2.9, 11.4, 1.2, subtitle, 18, theme["ink"])
        return
    if motif == "poster":
        _rect(slide, inches, 0, 0, 13.333, 7.5, theme["paper"])
        for index, key in enumerate(("accent", "gold", "navy2", "navy")):
            _rect(slide, inches, index * 3.333, 0, 3.333, 1.15, theme[key])
        _tb(slide, inches, pt, 0.7, 2.0, 11.6, 1.6, title, 36, theme["ink"], bold=True)
        if subtitle:
            _tb(slide, inches, pt, 0.7, 3.8, 11.4, 1.1, subtitle, 18, theme["slate"])
        return
    _rect(slide, inches, 0, 0, 13.333, 7.5, theme["paper"])
    if motif == "dashboard":
        _rect(slide, inches, 0, 0, 13.333, 2.05, theme["white"])
        _tb(slide, inches, pt, 0.55, 0.7, 11.5, 0.9, title, 32, theme["ink"], bold=True)
    elif motif in ("consulting", "ink", "sketch"):
        _tb(slide, inches, pt, 0.7, 1.9, 11.6, 1.8, title, 40, theme["ink"], bold=True)
        _rect(slide, inches, 0.7, 4.0, 2.2, 0.03, theme["accent"])
    else:
        _rect(slide, inches, 0, 0, 0.2, 7.5, theme["accent"])
        _tb(slide, inches, pt, 0.7, 1.9, 11.6, 1.8, title, 40, theme["ink"], bold=True)
    if subtitle and motif != "dashboard":
        _tb(slide, inches, pt, 0.7, 4.3, 11.4, 1.1, subtitle, 18, theme["slate"])
    elif subtitle:
        _tb(slide, inches, pt, 0.55, 2.4, 12.0, 0.8, subtitle, 16, theme["slate"])
    _page_mark(slide, inches, pt, 1, total, theme["mute"])


def _paint_section(slide, inches, pt, motif, theme, kicker, title, bullets, page, total):
    if motif == "official":
        _rect(slide, inches, 0, 0, 13.333, 7.5, theme["accent"])
        title_color, sub_color = theme["white"], theme["white"]
    elif motif in ("defense", "lesson"):
        _rect(slide, inches, 0, 0, 13.333, 7.5, theme["navy"])
        title_color, sub_color = theme["white"], theme["white"]
    elif motif == "split":
        _rect(slide, inches, 0, 0, 7.4, 7.5, theme["navy"])
        _rect(slide, inches, 7.4, 0, 5.933, 7.5, theme["paper"])
        _rect(slide, inches, 7.4, 0, 0.1, 7.5, theme["accent"])
        _tb(slide, inches, pt, 0.55, 2.5, 6.4, 1.6, title, 32, theme["white"], bold=True)
        _page_mark(slide, inches, pt, page, total, theme["mute"])
        return
    else:
        _rect(slide, inches, 0, 0, 13.333, 7.5, theme["paper"])
        _rect(slide, inches, 0.75, 4.2, 1.6, 0.035, theme["accent"])
        title_color, sub_color = theme["navy"], theme["slate"]
    if kicker:
        kick = theme["white"] if motif in ("official", "defense", "lesson") else theme["accent"]
        _tb(slide, inches, pt, 0.75, 2.05, 11.5, 0.32, kicker, 14, kick, bold=True)
    _tb(slide, inches, pt, 0.75, 2.45, 11.6, 1.35, title, 36, title_color, bold=True)
    _lines(slide, inches, pt, 0.75, 4.4, 11.2, 1.6, bullets, 18, sub_color)
    _page_mark(slide, inches, pt, page, total, sub_color)


def _split_card(text):
    for sep in ("：", ": ", " — ", " – ", " - ", " | "):
        if sep in text:
            head, body = text.split(sep, 1)
            head, body = head.strip(), body.strip()
            if head and body and len(head) <= 16:
                return head, body
    return "", text


def _split_kpi(text):
    for sep in (" | ", "｜", " — ", " – "):
        if sep in text:
            value, label = text.split(sep, 1)
            value, label = value.strip(), label.strip()
            if value and label:
                return value, label
    return text, ""


def _cards(slide, inches, pt, theme, bullets, top, motif="report"):
    texts = _texts(bullets)
    n = len(texts)
    cols, rows = (2, 2) if n == 4 else ((n, 1) if n <= 3 else (3, 2))
    gap = 0.16
    usable_w, usable_h = 12.43, 6.90 - top
    card_w = (usable_w - gap * (cols - 1)) / cols
    card_h = (usable_h - gap * (rows - 1)) / rows
    card_h = min(card_h, 2.70 if rows == 1 else 2.40)
    ladder = [theme["accent"], theme["navy2"], theme["gold"], theme["navy"]]
    for index, text in enumerate(texts):
        col, row = index % cols, index // cols
        x = 0.48 + col * (card_w + gap)
        y = top + row * (card_h + gap)
        cap = ladder[index % len(ladder)]
        if motif == "defense":
            cap = theme["accent"] if (n > 1 and index == n - 1) else theme["navy"]
        elif motif == "official":
            cap = theme["accent"]
        rounded = motif in ("craft", "board", "lesson")
        _rect(slide, inches, x, y, card_w, card_h, theme["card"], rounded=rounded)
        if motif == "split":
            _rect(slide, inches, x, y, 0.10, card_h, cap)
        elif motif not in ("ink", "consulting"):
            _rect(slide, inches, x, y, card_w, 0.075, cap)
            if motif == "official":
                _rect(slide, inches, x, y + 0.075, card_w, 0.025, theme["gold"])
        else:
            _rect(slide, inches, x, y + card_h - 0.02, card_w, 0.02, theme["accent"])
        _tb(slide, inches, pt, x + 0.18, y + 0.18, card_w - 0.36, 0.32, f"{index + 1:02d}", 12, cap, bold=True)
        head, body = _split_card(text)
        if head:
            _tb(slide, inches, pt, x + 0.18, y + 0.52, card_w - 0.36, 0.6, head, 16, theme["navy"], bold=True)
            _tb(slide, inches, pt, x + 0.18, y + 1.14, card_w - 0.36, card_h - 1.35, body, 14, theme["slate"])
        else:
            _tb(slide, inches, pt, x + 0.18, y + 0.52, card_w - 0.36, card_h - 0.75, text, 15, theme["ink"])


def _timeline(slide, inches, pt, theme, bullets, top):
    texts = _texts(bullets)
    n = len(texts)
    cols = n if n <= 4 else 3
    rows = 1 if n <= 4 else 2
    gap = 0.18
    usable_w, usable_h = 12.43, 6.90 - top
    card_w = (usable_w - gap * (cols - 1)) / cols
    row_h = min(3.10 if rows == 1 else 2.20, (usable_h - gap * (rows - 1)) / rows)
    ladder = [theme["accent"], theme["navy2"], theme["gold"], theme["navy"]]
    for r in range(rows):
        ry = top + r * (row_h + gap)
        rc = cols if r < rows - 1 else min(cols, n - r * cols)
        if rc > 1:
            _rect(slide, inches, 0.48 + card_w / 2, ry + 0.20, (rc - 1) * (card_w + gap), 0.04, theme["accent"])
    for index, text in enumerate(texts):
        col, row = index % cols, index // cols
        x = 0.48 + col * (card_w + gap)
        ry = top + row * (row_h + gap)
        cap = ladder[index % len(ladder)]
        _rect(slide, inches, x + (card_w - 0.46) / 2, ry, 0.46, 0.44, cap, rounded=True)
        _tb(slide, inches, pt, x + (card_w - 0.46) / 2, ry + 0.04, 0.46, 0.36, f"{index + 1:02d}", 12, theme["white"], bold=True, align="center")
        cy = ry + 0.58
        ch = max(0.9, row_h - 0.58)
        _rect(slide, inches, x, cy, card_w, ch, theme["card"], rounded=True)
        _rect(slide, inches, x, cy, card_w, 0.07, cap)
        head, body = _split_card(text)
        if head:
            _tb(slide, inches, pt, x + 0.16, cy + 0.16, card_w - 0.32, 0.48, head, 15, theme["navy"], bold=True)
            _tb(slide, inches, pt, x + 0.16, cy + 0.66, card_w - 0.32, ch - 0.76, body, 13, theme["slate"])
        else:
            _tb(slide, inches, pt, x + 0.16, cy + 0.18, card_w - 0.32, ch - 0.30, text, 14, theme["ink"])


def _comparison(slide, inches, pt, theme, bullets, top):
    texts = _texts(bullets)
    n = len(texts)
    col_gap = 0.34
    col_w = (12.43 - col_gap) / 2
    rows = (n + 1) // 2
    gap = 0.16
    card_h = min(3.20 if rows == 1 else 1.90, (6.90 - top - gap * (rows - 1)) / rows)
    _rect(slide, inches, 0.48 + col_w + col_gap / 2 - 0.01, top, 0.02, rows * card_h + (rows - 1) * gap, theme["accent"])
    for index, text in enumerate(texts):
        col, row = index % 2, index // 2
        x = 0.48 + col * (col_w + col_gap)
        y = top + row * (card_h + gap)
        side = theme["navy"] if col == 0 else theme["accent"]
        _rect(slide, inches, x, y, col_w, card_h, theme["card"], rounded=True)
        _rect(slide, inches, x, y, 0.10, card_h, side)
        _tb(slide, inches, pt, x + 0.20, y + 0.14, col_w - 0.36, 0.28, f"{index + 1:02d}", 12, side, bold=True)
        head, body = _split_card(text)
        if head:
            _tb(slide, inches, pt, x + 0.20, y + 0.44, col_w - 0.36, 0.46, head, 16, side, bold=True)
            _tb(slide, inches, pt, x + 0.20, y + 0.94, col_w - 0.36, card_h - 1.06, body, 14, theme["slate"])
        else:
            _tb(slide, inches, pt, x + 0.20, y + 0.46, col_w - 0.36, card_h - 0.60, text, 15, theme["ink"])


def _agenda(slide, inches, pt, theme, bullets, top):
    texts = _texts(bullets)
    n = len(texts)
    gap = 0.12
    row_h = min((6.90 - top - gap * (n - 1)) / n, 0.92)
    for index, text in enumerate(texts):
        y = top + index * (row_h + gap)
        _rect(slide, inches, 0.48, y, 12.43, row_h, theme["card"], rounded=True)
        _rect(slide, inches, 0.64, y + (row_h - 0.42) / 2, 0.72, 0.42, theme["accent"], rounded=True)
        _tb(slide, inches, pt, 0.64, y + (row_h - 0.42) / 2, 0.72, 0.42, f"{index + 1:02d}", 14, theme["white"], bold=True, align="center")
        _tb(slide, inches, pt, 1.56, y + (row_h - 0.42) / 2, 11.0, 0.42, text, 16, theme["ink"])


def _kpi(slide, inches, pt, theme, bullets, top):
    texts = _texts(bullets)
    n = len(texts)
    cols = n if n <= 4 else 3
    rows = 1 if n <= 4 else 2
    gap = 0.16
    card_w = (12.43 - gap * (cols - 1)) / cols
    card_h = min(2.15, (6.90 - top - gap * (rows - 1)) / rows)
    for index, text in enumerate(texts):
        col, row = index % cols, index // cols
        x = 0.48 + col * (card_w + gap)
        y = top + row * (card_h + gap)
        _rect(slide, inches, x, y, card_w, card_h, theme["card"], rounded=True)
        _rect(slide, inches, x, y, card_w, 0.07, theme["accent"])
        value, label = _split_kpi(text)
        _tb(slide, inches, pt, x + 0.18, y + 0.28, card_w - 0.36, 0.7, value, 22 if cols >= 4 else 26, theme["navy"], bold=True)
        if label:
            _tb(slide, inches, pt, x + 0.18, y + 1.05, card_w - 0.36, 0.7, label, 13, theme["slate"])


def _bullets(slide, inches, pt, theme, bullets, top, bottom=6.90):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    texts = _texts(bullets)
    if not texts or bottom - top < 0.4:
        return
    box = slide.shapes.add_textbox(inches(0.48), inches(top), inches(12.43), inches(bottom - top))
    frame = box.text_frame
    frame.word_wrap = True
    for index, text in enumerate(texts):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(10)
        run = paragraph.add_run()
        run.text = "•  " + text
        _set_run_font(run, size=pt(18), color=theme["ink"])


def _images(slide, inches, item, top):
    images = item.get("images") or []
    if isinstance(images, dict):
        images = [images]
    paths = []
    for image in images:
        if isinstance(image, str):
            path = image
        elif isinstance(image, dict):
            path = image.get("path") or ""
        else:
            continue
        path = str(path).strip()
        if not path:
            continue
        if not os.path.isfile(path):
            print(f"ERROR: slide image not found: {path}", file=sys.stderr)
            sys.exit(1)
        paths.append(path)
    if not paths:
        return top
    slot = 12.43 / len(paths)
    y = 6.90 - 2.3
    if y < top + 0.4:
        y = top
    for index, path in enumerate(paths):
        slide.shapes.add_picture(path, inches(0.48 + index * slot + 0.08), inches(y), height=inches(2.05))
    return y - 0.1


def _chart_type(name, enum):
    key = str(name or "").strip().lower().replace("-", " ").replace("_", " ")
    if key in ("bar h", "bar horizontal", "horizontal bar", "条形"):
        return enum.BAR_CLUSTERED
    if key in ("line", "折线"):
        return enum.LINE
    if key in ("pie", "饼"):
        return enum.PIE
    if key in ("area", "面积"):
        return enum.AREA
    if key in ("radar", "雷达") and hasattr(enum, "RADAR"):
        return enum.RADAR
    return enum.COLUMN_CLUSTERED


def _charts(slide, inches, item, top, bottom=6.85):
    charts = item.get("charts") or []
    if not charts:
        return
    try:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
    except ImportError:
        print("ERROR: python-pptx chart support is unavailable.", file=sys.stderr)
        sys.exit(1)
    slot = 12.43 / len(charts)
    for index, spec in enumerate(charts):
        if not isinstance(spec, dict):
            continue
        data = CategoryChartData()
        data.categories = [_clean_xml_text(c) for c in (spec.get("categories") or [])]
        for series in spec.get("series") or []:
            if isinstance(series, dict):
                data.add_series(_clean_xml_text(series.get("name") or "系列"), series.get("values") or [])
        graphic = slide.shapes.add_chart(
            _chart_type(spec.get("chart_type"), XL_CHART_TYPE),
            inches(0.48 + index * slot), inches(top), inches(max(1.2, slot - 0.16)), inches(max(1.6, bottom - top)),
            data,
        )
        chart = graphic.chart
        chart.has_legend = len(spec.get("series") or []) > 1
        title = _clean_xml_text(spec.get("title") or "")
        chart.has_title = bool(title)
        if title:
            chart.chart_title.text_frame.paragraphs[0].text = title


def _add_notes(slide, notes):
    notes = _clean_xml_text(notes)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def build_pptx(outline, output_path):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        print("ERROR: python-pptx not installed. Run: pip install python-pptx", file=sys.stderr)
        sys.exit(1)

    style = _resolve_theme(outline.get("theme"), _style_hint(outline))
    motif = style.get("motif") or "report"
    theme = {key: _rgb(value) for key, value in style["colors"].items()}
    title = _clean_xml_text(outline.get("title") or "演示文稿")
    subtitle = _clean_xml_text(outline.get("subtitle") or "")
    slides = [item for item in outline["slides"] if isinstance(item, dict)]
    footer = title if len(title) <= 36 else title[:35] + "…"
    total = len(slides) + 1

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = title
    blank = _blank_layout(prs)

    cover = prs.slides.add_slide(blank)
    _paint_cover(cover, Inches, Pt, motif, theme, title, subtitle, total)

    for index, item in enumerate(slides, start=2):
        slide = prs.slides.add_slide(blank)
        kind = _layout(item)
        slide_title = _clean_xml_text(item.get("title") or f"第 {index - 1} 页")
        kicker = _clean_xml_text(item.get("kicker") or "")
        bullets = item.get("bullets")
        if kind in ("section", "closing"):
            _paint_section(slide, Inches, Pt, motif, theme, kicker, slide_title, bullets, index, total)
        elif kind == "quote":
            texts = _texts(bullets)
            label = kicker or (slide_title if texts else "")
            quote = texts or [slide_title]
            top = _light_chrome(slide, Inches, Pt, theme, label, "", footer, index, total, motif)
            _tb(slide, Inches, Pt, 0.48, top, 1.4, 0.8, "“", 60, theme["accent"], bold=True)
            _lines(slide, Inches, Pt, 0.48, top + 0.85, 12.2, 4.6, quote, 26, theme["navy"], space=10)
        else:
            top = _light_chrome(slide, Inches, Pt, theme, kicker, slide_title, footer, index, total, motif)
            if kind == "cards":
                _cards(slide, Inches, Pt, theme, bullets, top, motif)
            elif kind == "timeline":
                _timeline(slide, Inches, Pt, theme, bullets, top)
            elif kind == "comparison":
                _comparison(slide, Inches, Pt, theme, bullets, top)
            elif kind == "agenda":
                _agenda(slide, Inches, Pt, theme, bullets, top)
            elif kind == "kpi":
                _kpi(slide, Inches, Pt, theme, bullets, top)
            else:
                has_images = bool(item.get("images"))
                has_charts = bool(item.get("charts"))
                bullet_bottom = 6.90
                if has_images:
                    bullet_bottom = min(bullet_bottom, 4.50)
                    _images(slide, Inches, item, top)
                if has_charts:
                    chart_top = top if not _texts(bullets) else min(top + 1.7, 4.3)
                    bullet_bottom = min(bullet_bottom, chart_top - 0.08)
                    _charts(slide, Inches, item, chart_top, 4.40 if has_images else 6.85)
                _bullets(slide, Inches, Pt, theme, bullets, top, bullet_bottom)
        _add_notes(slide, item.get("notes") or "")

    prs.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Build a designed .pptx deck from a JSON outline")
    parser.add_argument("--input", default="", help="Outline JSON file path (omit to read from stdin)")
    parser.add_argument("--output", default=".", help="Output .pptx file path or directory")
    args = parser.parse_args()

    outline = _load_outline(args)
    output_path = _resolve_output_path(args.output, outline.get("title"))
    result = build_pptx(outline, output_path)
    print(f"PPTX 生成完成: {result}")
    print(f"   页数: {len(outline['slides']) + 1}")
    print(f"   输出: {os.path.abspath(result)}")


if __name__ == "__main__":
    main()
